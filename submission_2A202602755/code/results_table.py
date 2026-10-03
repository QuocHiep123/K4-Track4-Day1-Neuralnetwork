"""results_table.py — Lưu kết quả JSON và tự động điền bảng Excel.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (giữ nguyên công thức và 4 sheet).
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result['cfg'], result['history'], result['summary'] ra <results_dir>/<exp_id>.json."""
    p = Path(results_dir)
    p.mkdir(parents=True, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    file_path = p / f"{exp_id}.json"

    data_to_save = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result.get("summary", {}),
    }

    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    return str(file_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    p = Path(results_dir)
    if not p.exists():
        return []

    results = []
    for f in sorted(p.glob("*.json")):
        with open(f, "r", encoding="utf-8") as fp:
            results.append(json.load(fp))

    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng Experiments."""
    cfg = result["cfg"]
    summary = result.get("summary", {})
    exp_id = cfg.get("exp_id", "")

    hidden = cfg.get("hidden", (256, 128))
    if isinstance(hidden, (list, tuple)):
        hidden_str = "-".join(str(h) for h in hidden)
    else:
        hidden_str = str(hidden)

    clip_norm = cfg.get("clip_norm")
    clip_val = "none" if clip_norm is None else clip_norm

    diverged = summary.get("diverged", False)

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce").upper(),
        "optimizer": cfg.get("optimizer", ""),
        "lr": cfg.get("lr"),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": hidden_str,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": clip_val,
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss"),
        "best_val_loss": summary.get("best_val_loss"),
        "best_epoch": summary.get("best_epoch"),
        "final_train_loss": summary.get("final_train_loss"),
        "final_val_loss": summary.get("final_val_loss"),
        "val_acc": summary.get("val_acc"),
        "val_macro_f1": summary.get("val_macro_f1"),
        "time_per_epoch_s": summary.get("time_per_epoch_s"),
        "peak_mem_MB": summary.get("peak_mem_MB"),
        "diverged": "Có" if diverged else "Không",
        "eval_acc": eval_scores.get("accuracy") if eval_scores else None,
        "eval_macro_f1": eval_scores.get("macro_f1") if eval_scores else None,
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet 'Experiments' của mẫu, bảo tồn công thức và lưu ra out_path."""
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    header_cols = {}
    for col_idx in range(1, ws.max_column + 1):
        name = ws.cell(1, col_idx).value
        if name:
            header_cols[name] = col_idx

    formula_names = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}

    for idx, r_data in enumerate(rows):
        row_num = idx + 2
        for key, val in r_data.items():
            if key in header_cols and key not in formula_names:
                ws.cell(row_num, header_cols[key], val)

        # Ghi các công thức tự động cho mỗi dòng
        # P = col 16 (step0_loss), S = col 19 (final_train_loss), T = col 20 (final_val_loss), V = col 22 (val_macro_f1), AF = col 32 (delta_val_f1_vs_base)
        if "step0_gap_vs_lnC" in header_cols:
            ws.cell(row_num, header_cols["step0_gap_vs_lnC"], f'=IF(P{row_num}="","",P{row_num}-LN(7))')
        if "gap_val_minus_train" in header_cols:
            ws.cell(row_num, header_cols["gap_val_minus_train"], f'=IF(OR(T{row_num}="",S{row_num}=""),"",T{row_num}-S{row_num})')
        if "delta_val_f1_vs_base" in header_cols:
            ws.cell(row_num, header_cols["delta_val_f1_vs_base"], f'=IF(OR(V{row_num}="",Seeds!$C$8=""),"",V{row_num}-Seeds!$C$8)')
        if "beyond_noise" in header_cols:
            ws.cell(row_num, header_cols["beyond_noise"], f'=IF(OR(AF{row_num}="",Seeds!$C$10=""),"",IF(ABS(AF{row_num})>Seeds!$C$10,"Có","Không"))')

    out_p = Path(out_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Đã lưu bảng kết quả vào {out_path} ({len(rows)} thí nghiệm)")
